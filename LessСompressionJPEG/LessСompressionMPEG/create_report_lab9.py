from __future__ import annotations

from pathlib import Path

import pythoncom
import win32com.client
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


REPO_ROOT = Path(r"D:\Denys-526")
REPORT_DIR = REPO_ROOT / "screen"
DOCX_PATH = REPORT_DIR / "lab9.docx"
PDF_PATH = REPORT_DIR / "lab9.pdf"

RESULTS_DIR = next(REPO_ROOT.glob("Less*JPEG/Less*MPEG/Results"))

FIGURES = [
    (
        "First frame.png",
        "На рисунку 1 наведено опорний кадр, який використовується як I-кадр для подальшого прогнозування наступного кадру.",
        "Рисунок 1 – Перший (опорний) кадр відеопослідовності.",
    ),
    (
        "Second frame.png",
        "На рисунку 2 наведено поточний кадр, який підлягає кодуванню. Саме для нього виконується пошук подібних блоків у попередньому кадрі.",
        "Рисунок 2 – Другий (цільовий) кадр відеопослідовності.",
    ),
    (
        "Difference between frame.png",
        "На рисунку 3 показано міжкадрову різницю без компенсації руху. Більша частина кадру має малу інтенсивність змін, а локальні світлі ділянки відповідають зонам руху.",
        "Рисунок 3 – Абсолютна різниця між сусідніми кадрами.",
    ),
    (
        "Prediction frame.png",
        "На рисунку 4 наведено прогнозований кадр, який сформовано алгоритмом block matching із використанням блоку 16x16 та зони пошуку 7 пікселів.",
        "Рисунок 4 – Прогнозований кадр після блочного пошуку.",
    ),
    (
        "Residual frame.png",
        "На рисунку 5 наведено залишковий кадр. Низька енергія цього зображення підтверджує, що після компенсації руху для передавання залишається значно менше інформації.",
        "Рисунок 5 – Залишковий кадр після компенсації руху.",
    ),
    (
        "Restore frame.png",
        "На рисунку 6 наведено результат декодування. Відновлений кадр збігається з початковим цільовим кадром, що підтверджує коректність реалізації процедур обчислення залишку та реконструкції.",
        "Рисунок 6 – Відновлений кадр після декодування.",
    ),
    (
        "Гістограма кількості біт на піксель для різних варіантів кодування.png",
        "На рисунку 7 наведено порівняння кількості біт на піксель для початкового кадру, простої різниці між кадрами та залишкового сигналу після компенсації руху. Найменше значення отримано для залишкового кадру, що відповідає суті MPEG/H.264-подібного міжкадрового кодування.",
        "Рисунок 7 – Гістограма кількості біт на піксель для різних варіантів кодування.",
    ),
]


def set_run_style(run, size=14, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold


def add_paragraph(document: Document, text: str, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, first_line_cm=1.25, size=14):
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.first_line_indent = Cm(first_line_cm) if first_line_cm else None
    run = paragraph.add_run(text)
    set_run_style(run, size=size, bold=bold)
    return paragraph


def add_blank(document: Document):
    document.add_paragraph()


def ensure_inputs():
    REPORT_DIR.mkdir(exist_ok=True)
    for filename, _, _ in FIGURES:
        path = RESULTS_DIR / filename
        if not path.exists():
            raise FileNotFoundError(path)


def build_docx():
    ensure_inputs()
    document = Document()

    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)

    title_lines = [
        ("МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ", True),
        ("Національний аерокосмічний університет", False),
        ("«Харківський авіаційний інститут»", False),
        ("Факультет радіотехніки, комп’ютерних систем і інфокомунікацій", False),
        ("Кафедра інформаційно-комунікаційних технологій ім. О. О. Зеленського", False),
        ("", False),
        ("Лабораторна робота №9", True),
        ("з дисципліни «Основи теорії цифрового зв’язку»", False),
        ("", False),
        (
            "на тему: «Реалізація алгоритму H.264 (MPEG-4 Part 10) для стиснення "
            "відеозображень за допомогою мови програмування Python»",
            False,
        ),
        ("", False),
        ("", False),
        ("Виконав: студент 3 курсу групи № 526-СТ", False),
        ("напряму підготовки (спеціальності)", False),
        ("172 Телекомунікації та радіотехніка", False),
        ("Чернецький Денис Олександрович", False),
        ("", False),
        ("Прийняв: доц. В'юницький О.Г.", False),
        ("Національна шкала: __________", False),
        ("Кількість балів: _____", False),
        ("Оцінка: ECTS _______", False),
    ]

    for text, bold in title_lines:
        if not text:
            add_blank(document)
            continue
        align = WD_ALIGN_PARAGRAPH.CENTER if "Виконав" not in text and "Прийняв" not in text and "Національна" not in text and "Кількість" not in text and "Оцінка" not in text and "напряму" not in text and "172 " not in text and "Чернецький" not in text else WD_ALIGN_PARAGRAPH.LEFT
        add_paragraph(document, text, align=align, bold=bold, first_line_cm=0, size=14)

    while len(document.paragraphs) < 28:
        add_blank(document)

    add_paragraph(document, "Харків – 2026", align=WD_ALIGN_PARAGRAPH.CENTER, first_line_cm=0, size=14)
    document.add_page_break()

    add_paragraph(document, "МЕТА РОБОТИ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, first_line_cm=0, size=14)
    add_blank(document)
    add_paragraph(
        document,
        "Ознайомитися з принципами міжкадрового кодування відеоданих, реалізувати базові етапи алгоритму H.264/MPEG-подібного стиснення засобами Python та оцінити ефективність компенсації руху за метрикою кількості біт на піксель.",
    )
    add_blank(document)

    add_paragraph(document, "ХІД РОБОТИ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, first_line_cm=0, size=14)
    add_blank(document)
    add_paragraph(
        document,
        "У ході виконання лабораторної роботи було реалізовано зчитування двох сусідніх кадрів з відеофайлу, сегментацію кадру на блоки розміром 16x16, побудову зони пошуку в опорному кадрі та пошук найбільш схожих блоків за критерієм MAD. На основі знайдених блоків сформовано прогнозований кадр, обчислено залишковий кадр та виконано реконструкцію цільового кадру.",
    )
    add_paragraph(
        document,
        "Для оцінювання ефективності міжкадрового кодування було порівняно три варіанти представлення даних: початковий кадр, абсолютну різницю між сусідніми кадрами та залишковий кадр після компенсації руху. Для кожного з них розраховано середню кількість біт на піксель окремо для компонент R, G, B та для сумарного RGB-подання.",
    )
    add_blank(document)

    add_paragraph(document, "ОСНОВНІ РЕЗУЛЬТАТИ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, first_line_cm=0, size=14)
    add_blank(document)
    add_paragraph(
        document,
        "Для демонстрації роботи алгоритму було використано сусідні кадри №2585 та №2586 з відеофайлу sample4.avi. Під час перевірки встановлено, що реконструйований кадр повністю збігається з цільовим кадром, а коефіцієнт стиску за оцінкою bits/pixel становить 75.8199.",
    )
    add_paragraph(
        document,
        "Отримані числові результати мають такий вигляд: пара кадрів - 2585 та 2586; відновлення кадру - повний збіг із цільовим кадром; сумарне значення bits/pixel для початкового кадру RGB - 18.6566; для міжкадрової різниці - 0.6199; для залишкового кадру - 0.2461; коефіцієнт стиснення - 75.8199.",
    )
    add_paragraph(
        document,
        "Для окремих кольорових компонент отримано такі значення: для каналу R - 6.5216 / 0.2065 / 0.0811, для каналу G - 6.4736 / 0.2059 / 0.0820, для каналу B - 5.6614 / 0.2074 / 0.0830, де в кожній трійці наведено bits/pixel для початкового кадру, різниці між кадрами та залишкового кадру відповідно.",
    )
    add_blank(document)

    for filename, text, caption in FIGURES:
        add_paragraph(document, text)
        image_paragraph = document.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.add_run().add_picture(str(RESULTS_DIR / filename), width=Cm(16))
        add_paragraph(document, caption, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_cm=0, size=12)
        add_blank(document)

    add_paragraph(document, "ВИСНОВКИ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, first_line_cm=0, size=14)
    add_blank(document)
    add_paragraph(
        document,
        "У лабораторній роботі реалізовано базову схему міжкадрового відеокодування з компенсацією руху, яка включає пошук подібних блоків, побудову прогнозованого кадру, формування залишкового сигналу та реконструкцію зображення. Отримані результати показали, що після компенсації руху енергетика залишкового кадру істотно зменшується порівняно з початковим кадром і навіть зі звичайною різницею між кадрами.",
    )
    add_paragraph(
        document,
        "Для досліджуваної пари кадрів сумарна оцінка bits/pixel зменшилась з 18.6566 для початкового кадру до 0.2461 для залишкового кадру, що відповідає коефіцієнту стиску 75.8199. Відновлений кадр повністю збігається з цільовим, отже реалізація алгоритму кодування та декодування є коректною.",
    )

    if DOCX_PATH.exists():
        DOCX_PATH.unlink()
    document.save(str(DOCX_PATH))


def build_pdf():
    if PDF_PATH.exists():
        PDF_PATH.unlink()

    pythoncom.CoInitialize()
    word = None
    doc = None
    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(str(DOCX_PATH))
        doc.ExportAsFixedFormat(str(PDF_PATH), 17)
        doc.Close(False)
        doc = None
        word.Quit()
        word = None
    finally:
        if doc is not None:
            doc.Close(False)
        if word is not None:
            word.Quit()
        pythoncom.CoUninitialize()


if __name__ == "__main__":
    build_docx()
    build_pdf()
    print(f"Created: {DOCX_PATH}")
    print(f"Created: {PDF_PATH}")
